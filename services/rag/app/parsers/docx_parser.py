"""DOCX parser — olimpiyat formatı: bold + numaralı paragraf soru başlangıcı.

Bir soru bloğu:
  **3.** ... soru metni
     a) ...
     b) ...
     c) ...
     d) ...
     e) ...
"""
import logging
import re
from pathlib import Path
from typing import Iterator, Optional

from docx import Document
from docx.text.paragraph import Paragraph

from ..schemas import ParsedChunk
from .excel_parser import ExcelIndex

logger = logging.getLogger(__name__)


_LEADING_NUMBER_RE = re.compile(r"^\s*\*?\*?(\d{1,3})[.\)]\s*", re.MULTILINE)


def _is_bold_paragraph(p: Paragraph) -> bool:
    """Soru başlangıcı tespiti için "yeterince bold" kuralı.

    Olimpiyat formatında soru numarası bold, gövde normal — yani
    paragrafın çoğu bold olmuyor. Bu yüzden:
      a) İLK metin-run'u bold ise (numara bold) — kabul, VEYA
      b) Tüm metin-run'ları bold ise (tamamı bold paragraf) — kabul
    """
    runs = [r for r in p.runs if r.text and r.text.strip()]
    if not runs:
        return False
    # (a) ilk run bold
    if runs[0].bold:
        return True
    # (b) tüm runs bold
    return all(r.bold for r in runs)


def _leading_number(text: str) -> Optional[int]:
    m = _LEADING_NUMBER_RE.match(text)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            return None
    return None


class DocxParser:
    # ─── Soru havuzu ────────────────────────────────────────────────

    def parse_question_docx(
        self,
        file_path: str | Path,
        excel_meta: ExcelIndex,
        grade_level: Optional[int] = None,
        subject: Optional[str] = None,
    ) -> list[ParsedChunk]:
        path = Path(file_path)
        file_name = path.stem  # olimpiyat_2020.docx → olimpiyat_2020
        doc = Document(str(path))

        # Önce bütün paragrafları sırayla gez; bold + numaralı olanları
        # soru başlangıcı say. Aynı soruya ait blok bir sonraki bold
        # numaralı paragraf gelene kadar uzar.
        questions: list[tuple[int, list[str]]] = []
        current: list[str] | None = None
        current_qnum: int | None = None

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            qnum = _leading_number(text)
            if qnum is not None and _is_bold_paragraph(para):
                if current is not None and current_qnum is not None:
                    questions.append((current_qnum, current))
                current = [text]
                current_qnum = qnum
            else:
                if current is not None:
                    current.append(text)

        if current is not None and current_qnum is not None:
            questions.append((current_qnum, current))

        chunks: list[ParsedChunk] = []
        for qnum, lines in questions:
            body = "\n".join(lines).strip()
            if len(body) < 20:
                continue

            meta = excel_meta.get((file_name, qnum))
            if meta is None:
                taxonomic_level = "unknown"
                answer_key = None
                test_no = None
            else:
                taxonomic_level = meta.taxonomic_level
                answer_key = meta.answer_key
                test_no = meta.test_no

            text_content = f"Soru {qnum}: {body}"
            chunks.append(ParsedChunk(
                chunk_id=f"{file_name}#q{qnum}",
                text_content=text_content,
                text_for_embedding=text_content,
                has_image=False,
                metadata={
                    "content_type": "question",
                    "file_name": file_name,
                    "question_number": qnum,
                    "taxonomic_level": taxonomic_level,
                    "answer_key": answer_key,
                    "test_no": test_no,
                    # Admin upload formu birinci öncelikli; yoksa file_name fallback.
                    "subject": subject if subject else file_name,
                    "grade_level": grade_level,
                },
            ))

        logger.info("DOCX '%s' parsed: %d question chunks", file_name, len(chunks))
        return chunks

    # ─── Konu anlatımı ──────────────────────────────────────────────

    def parse_theory_docx(
        self,
        file_path: str | Path,
        grade_level: int,
        subject: str,
    ) -> list[ParsedChunk]:
        path = Path(file_path)
        file_name = path.stem
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        from ..chunker import Chunker
        ch = Chunker()
        sections = ch.split_by_headings(text)
        sections = [s for s in (s.strip() for s in sections) if s]
        chunks_text: list[str] = []
        for sec in sections:
            chunks_text.extend(ch.split_long_chunk(sec, max_tokens=500))
        chunks_text = ch.add_overlap(chunks_text, overlap_sentences=1)

        out: list[ParsedChunk] = []
        for i, body in enumerate(chunks_text):
            if len(body) < 30:
                continue
            out.append(ParsedChunk(
                chunk_id=f"{file_name}#t{i}",
                text_content=body,
                text_for_embedding=body,
                has_image=False,
                metadata={
                    "content_type": "theory",
                    "file_name": file_name,
                    "grade_level": grade_level,
                    "subject": subject,
                    "test_no": None,
                    "question_number": None,
                    "taxonomic_level": None,
                    "answer_key": None,
                },
            ))
        logger.info("Theory DOCX '%s': %d chunks", file_name, len(out))
        return out
